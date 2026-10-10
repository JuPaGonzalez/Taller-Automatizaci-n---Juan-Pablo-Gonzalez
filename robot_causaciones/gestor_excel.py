import re
from datetime import date, datetime
from pathlib import Path
from typing import Any
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

from modelo import FacturaEntrada, ResultadoFactura


def limpiar_nit(nit_raw: Any) -> int:
    """Limpia el NIT eliminando puntos, comas, espacios y guiones, retornándolo como entero."""
    if nit_raw is None:
        raise ValueError("El NIT no puede ser nulo")
    if isinstance(nit_raw, (int, float)):
        nit_str = str(int(nit_raw))
    else:
        nit_str = str(nit_raw).strip()
    # Eliminar cualquier carácter que no sea dígito
    solo_digitos = re.sub(r"\D", "", nit_str)
    if not solo_digitos:
        raise ValueError(f"NIT inválido: {nit_raw}")
    return int(solo_digitos)


def parsear_valor(valor_raw: Any) -> int | None:
    """Parsea el valor y verifica que sea un entero estrictamente mayor que cero."""
    if valor_raw is None:
        return None
    if isinstance(valor_raw, int):
        return valor_raw if valor_raw > 0 else None
    if isinstance(valor_raw, float):
        if valor_raw.is_integer() and valor_raw > 0:
            return int(valor_raw)
        return None

    # Si es cadena de texto
    v_str = str(valor_raw).strip()
    # Si contiene punto decimal con decimales no nulos, no es entero
    if "." in v_str and not re.match(r"^\d{1,3}(\.\d{3})*$", v_str):
        # Podría ser float tipo 100.5
        try:
            val_fl = float(v_str.replace(",", ""))
            if not val_fl.is_integer() or val_fl <= 0:
                return None
            return int(val_fl)
        except ValueError:
            return None

    # Quitar puntos o comas de miles
    limpio = re.sub(r"[,\s]", "", v_str)
    if "." in limpio:
        limpio = limpio.replace(".", "")
    if limpio.isdigit() and int(limpio) > 0:
        return int(limpio)
    return None


def parsear_fecha(fecha_raw: Any) -> tuple[date | None, str | None]:
    """
    Parsea la fecha de factura.
    Retorna (fecha, error_detalle).
    """
    if fecha_raw is None:
        return None, "La fecha de factura debe tener el formato dd/mm/aaaa"

    fecha_obj: date | None = None
    if isinstance(fecha_raw, datetime):
        fecha_obj = fecha_raw.date()
    elif isinstance(fecha_raw, date):
        fecha_obj = fecha_raw
    elif isinstance(fecha_raw, str):
        cadena = fecha_raw.strip()
        # Intentar formato yyyy-MM-dd del contrato de PorRadicar
        formatos = ["%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y"]
        for fmt in formatos:
            try:
                fecha_obj = datetime.strptime(cadena, fmt).date()
                break
            except ValueError:
                continue

    if fecha_obj is None:
        return None, "La fecha de factura debe tener el formato dd/mm/aaaa"

    if fecha_obj > date.today():
        return None, "La fecha de factura no puede ser posterior a la fecha actual"

    return fecha_obj, None


def leer_facturas_por_radicar(ruta: Path) -> list[tuple[FacturaEntrada | None, ResultadoFactura | None]]:
    """
    Lee robot_causaciones/datos/por_radicar.xlsx de la primera hoja.
    Realiza las validaciones de negocio previas a tocar SICOF.
    Retorna una lista de tuplas (FacturaEntrada, ResultadoFactura Rechazado inmediato).
    """
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []

    encabezados = [str(col).strip() if col is not None else "" for col in filas[0]]
    idx_id = encabezados.index("IdLista") if "IdLista" in encabezados else 0
    idx_fac = encabezados.index("Factura") if "Factura" in encabezados else 1
    idx_nit = encabezados.index("NIT") if "NIT" in encabezados else 2
    idx_prov = encabezados.index("Proveedor") if "Proveedor" in encabezados else 3
    idx_val = encabezados.index("Valor") if "Valor" in encabezados else 4
    idx_centro = encabezados.index("CentroCosto") if "CentroCosto" in encabezados else 5
    idx_fec = encabezados.index("FechaFactura") if "FechaFactura" in encabezados else 6

    elementos: list[tuple[FacturaEntrada | None, ResultadoFactura | None]] = []

    for fila in filas[1:]:
        if fila is None or len(fila) <= idx_id or fila[idx_id] is None or str(fila[idx_id]).strip() == "":
            continue

        try:
            id_lista = int(float(fila[idx_id]))
        except (ValueError, TypeError):
            continue

        factura_str = str(fila[idx_fac]).strip() if fila[idx_fac] is not None else ""
        proveedor_str = str(fila[idx_prov]).strip() if fila[idx_prov] is not None else ""
        centro_str = str(fila[idx_centro]).strip() if fila[idx_centro] is not None else ""

        # Limpiar NIT
        try:
            nit_limpio = limpiar_nit(fila[idx_nit])
        except Exception:
            elementos.append((
                None,
                ResultadoFactura(
                    id_lista=id_lista,
                    factura=factura_str,
                    nit=0,
                    estado="RECHAZADA",
                    detalle="El NIT no es válido",
                ),
            ))
            continue

        # Validación previa 1: Factura obligatoria
        if not factura_str:
            elementos.append((
                None,
                ResultadoFactura(
                    id_lista=id_lista,
                    factura=factura_str,
                    nit=nit_limpio,
                    estado="RECHAZADA",
                    detalle="El número de factura es obligatorio",
                ),
            ))
            continue

        # Validación previa 2: Valor entero > 0
        valor_valido = parsear_valor(fila[idx_val])
        if valor_valido is None:
            elementos.append((
                None,
                ResultadoFactura(
                    id_lista=id_lista,
                    factura=factura_str,
                    nit=nit_limpio,
                    estado="RECHAZADA",
                    detalle="El valor debe ser un número entero mayor que cero",
                ),
            ))
            continue

        # Validación previa 3: Fecha válida y no posterior a hoy
        fecha_obj, error_fecha = parsear_fecha(fila[idx_fec])
        if error_fecha:
            elementos.append((
                None,
                ResultadoFactura(
                    id_lista=id_lista,
                    factura=factura_str,
                    nit=nit_limpio,
                    estado="RECHAZADA",
                    detalle=error_fecha,
                ),
            ))
            continue

        factura_entrada = FacturaEntrada(
            id_lista=id_lista,
            factura=factura_str,
            nit=nit_limpio,
            proveedor=proveedor_str,
            valor=valor_valido,
            centro_costo=centro_str,
            fecha_factura=fecha_obj,  # type: ignore
        )
        elementos.append((factura_entrada, None))

    return elementos


def guardar_resultados(resultados: list[ResultadoFactura], ruta: Path) -> None:
    """
    Escribe salida/resultados.xlsx con la tabla 'Resultados'
    según el contrato estricto de archivos del taller.
    """
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultados"

    encabezados = ["IdLista", "Factura", "NIT", "Estado", "Causacion", "Detalle"]
    ws.append(encabezados)

    for r in resultados:
        ws.append([
            int(r.id_lista),
            str(r.factura),
            int(r.nit),
            str(r.estado),
            str(r.causacion) if r.causacion else "",
            str(r.detalle) if r.detalle else "",
        ])

    num_filas = len(resultados) + 1
    ref = f"A1:F{max(num_filas, 2)}"
    tabla = Table(displayName="Resultados", ref=ref)
    estilo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    tabla.tableStyleInfo = estilo
    ws.add_table(tabla)

    wb.save(ruta)
